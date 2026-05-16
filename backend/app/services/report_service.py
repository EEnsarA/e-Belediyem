import io
import logging
from typing import Optional, List, Dict, Any
from datetime import datetime

logger = logging.getLogger(__name__)


class ReportService:
    """PDF ve Word rapor üretimi."""

    async def generate_pdf_report(
        self,
        municipality_name: str,
        stats: Dict[str, Any],
        complaints: List[Dict],
        ai_summary: Optional[str] = None,
        start_date: Optional[datetime] = None,
        end_date: Optional[datetime] = None,
    ) -> bytes:
        """ReportLab ile PDF rapor üret."""
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import cm
            from reportlab.lib import colors
            from reportlab.platypus import (
                SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
            )
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont

            buffer = io.BytesIO()
            doc = SimpleDocTemplate(
                buffer,
                pagesize=A4,
                rightMargin=2*cm,
                leftMargin=2*cm,
                topMargin=2*cm,
                bottomMargin=2*cm,
            )

            styles = getSampleStyleSheet()
            story = []

            # Başlık
            title_style = ParagraphStyle(
                "Title",
                parent=styles["Title"],
                fontSize=18,
                textColor=colors.HexColor("#1e40af"),
                spaceAfter=12,
            )
            story.append(Paragraph(f"{municipality_name}", title_style))
            story.append(Paragraph("Belediye Analiz ve Durum Raporu", styles["Heading2"]))

            date_range = ""
            if start_date and end_date:
                date_range = f"{start_date.strftime('%d.%m.%Y')} - {end_date.strftime('%d.%m.%Y')}"
            story.append(Paragraph(f"Dönem: {date_range or 'Tüm Zamanlar'}", styles["Normal"]))
            story.append(Paragraph(f"Oluşturulma: {datetime.now().strftime('%d.%m.%Y %H:%M')}", styles["Normal"]))
            story.append(Spacer(1, 0.5*cm))
            story.append(HRFlowable(width="100%", thickness=2, color=colors.HexColor("#1e40af")))
            story.append(Spacer(1, 0.5*cm))

            # Özet istatistikler tablosu
            story.append(Paragraph("Özet İstatistikler", styles["Heading2"]))
            data = [
                ["Metrik", "Değer"],
                ["Toplam Şikayet", str(stats.get("total", 0))],
                ["Çözülen", str(stats.get("resolved", 0))],
                ["Bekleyen", str(stats.get("pending", 0))],
                ["Çözüm Oranı", f"{stats.get('resolution_rate', 0):.1f}%"],
                ["Acil Şikayet", str(stats.get("urgent", 0))],
                ["Ort. Memnuniyet", f"{stats.get('avg_satisfaction', 'N/A')}"],
            ]
            table = Table(data, colWidths=[8*cm, 8*cm])
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e40af")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f0f9ff")]),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
                ("FONTSIZE", (0, 0), (-1, -1), 11),
                ("PADDING", (0, 0), (-1, -1), 6),
            ]))
            story.append(table)
            story.append(Spacer(1, 0.5*cm))

            # AI Özeti
            if ai_summary:
                story.append(Paragraph("AI Yönetici Analizi", styles["Heading2"]))
                story.append(Paragraph(ai_summary, styles["Normal"]))
                story.append(Spacer(1, 0.5*cm))

            # Son şikayetler
            if complaints:
                story.append(Paragraph("Şikayet Listesi", styles["Heading2"]))
                comp_data = [["#", "Açıklama", "Kategori", "Durum", "Tarih"]]
                for c in complaints[:20]:  # Max 20
                    comp_data.append([
                        str(c.get("id", "")),
                        (c.get("description", "")[:50] + "...") if len(c.get("description", "")) > 50 else c.get("description", ""),
                        c.get("category", "Diğer") or "Diğer",
                        c.get("status", ""),
                        c.get("created_at", ""),
                    ])
                comp_table = Table(comp_data, colWidths=[1.5*cm, 6*cm, 3*cm, 3*cm, 2.5*cm])
                comp_table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#374151")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f9fafb")]),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e5e7eb")),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("PADDING", (0, 0), (-1, -1), 4),
                ]))
                story.append(comp_table)

            doc.build(story)
            return buffer.getvalue()

        except ImportError:
            logger.warning("reportlab yüklü değil")
            return b"PDF ozelligi aktif degil"
        except Exception as e:
            logger.error(f"PDF rapor hatası: {e}")
            raise

    async def generate_word_report(
        self,
        municipality_name: str,
        stats: Dict[str, Any],
        complaints: List[Dict],
        ai_summary: Optional[str] = None,
    ) -> bytes:
        """python-docx ile Word raporu üret."""
        try:
            from docx import Document
            from docx.shared import Pt, RGBColor, Inches
            from docx.enum.text import WD_ALIGN_PARAGRAPH

            doc = Document()

            # Başlık
            title = doc.add_heading(municipality_name, level=1)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER

            doc.add_heading("Şikayet Analiz Raporu", level=2)
            doc.add_paragraph(f"Oluşturulma Tarihi: {datetime.now().strftime('%d.%m.%Y %H:%M')}")
            doc.add_paragraph()

            # İstatistikler
            doc.add_heading("Özet İstatistikler", level=2)
            table = doc.add_table(rows=1, cols=2)
            table.style = "Table Grid"
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = "Metrik"
            hdr_cells[1].text = "Değer"

            rows_data = [
                ("Toplam Şikayet", str(stats.get("total", 0))),
                ("Çözülen", str(stats.get("resolved", 0))),
                ("Bekleyen", str(stats.get("pending", 0))),
                ("Çözüm Oranı", f"{stats.get('resolution_rate', 0):.1f}%"),
                ("Acil Şikayet", str(stats.get("urgent", 0))),
            ]
            for metric, value in rows_data:
                row_cells = table.add_row().cells
                row_cells[0].text = metric
                row_cells[1].text = value

            doc.add_paragraph()

            if ai_summary:
                doc.add_heading("AI Yönetici Analizi", level=2)
                doc.add_paragraph(ai_summary)
                doc.add_paragraph()

            buffer = io.BytesIO()
            doc.save(buffer)
            return buffer.getvalue()

        except ImportError:
            logger.warning("python-docx yüklü değil")
            return b"Word ozelligi aktif degil"
        except Exception as e:
            logger.error(f"Word rapor hatası: {e}")
            raise


report_service = ReportService()
